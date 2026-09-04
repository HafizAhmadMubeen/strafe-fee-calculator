r"""
fill_calculator_gui.py
==================================================================================
Point-and-click front end for fill_calculator.py.

    .venv\Scripts\python fill_calculator_gui.py      (or double-click "Run GUI.bat")

Three steps, each unlocks the next:
  1. Connect Edge   - opens Edge (via launch_edge_debug.bat) on the Testing
                      Calculator with the debugging port on, and checks it.
  2. Choose Excel   - pick the .xlsx test workbook.
  3. Run it         - fills the calculator row by row and writes the Fee Amount
                      back.  Pauses once at ~10% (opens the file in Excel to
                      check).  Stop any time; rerun continues from where it
                      stopped and re-does only rows whose inputs you changed.
"""

from __future__ import annotations

import os
import queue
import threading
import tkinter as tk
from tkinter import filedialog, ttk
from tkinter.scrolledtext import ScrolledText

import fill_calculator as fc

# --------------------------------------------------------------- appearance ----

BG        = "#EEF1F4"
CARD      = "#FFFFFF"
BORDER    = "#DDE1E6"
INK       = "#1F2933"
MUTE      = "#6B7480"
ACCENT    = "#2563EB"
ACCENT_DK = "#1D4ED8"
OK        = "#157A3B"
BAD       = "#B4232A"
HDR_BG    = "#1E293B"
HDR_SUB   = "#94A3B8"
WARN_BG   = "#FEF4D6"
WARN_BD   = "#E7A400"
WARN_INK  = "#8A5A00"
LOG_BG    = "#FbFcFd"
LOG_INK   = "#26313C"

FONT   = ("Segoe UI", 10)
FBOLD  = ("Segoe UI", 10, "bold")
FSMALL = ("Segoe UI", 9)
FH1    = ("Segoe UI", 16, "bold")
FSUB   = ("Segoe UI", 9)
FMONO  = ("Consolas", 9)

_BTN_KINDS = {
    "primary": (ACCENT, ACCENT_DK, "#FFFFFF"),
    "danger":  ("#D6323A", "#B4232A", "#FFFFFF"),
    "ghost":   ("#E7EBEF", "#DCE1E6", INK),
    "go":      (OK, "#0F5F2D", "#FFFFFF"),
}


def make_button(parent, text, command, kind="primary", width=0):
    bg, abg, fg = _BTN_KINDS[kind]
    b = tk.Button(parent, text=text, command=command, font=FBOLD, fg=fg, bg=bg,
                  activebackground=abg, activeforeground=fg, relief="flat", bd=0,
                  padx=18, pady=9, cursor="hand2", highlightthickness=0,
                  disabledforeground="#9AA2AA")
    if width:
        b.configure(width=width)
    b._kind = kind
    return b


def set_button(b, enabled: bool):
    if enabled:
        bg, abg, fg = _BTN_KINDS[b._kind]
        b.configure(state="normal", bg=bg, activebackground=abg, fg=fg, cursor="hand2")
    else:
        b.configure(state="disabled", bg="#DDE2E7", cursor="arrow")


class Card(tk.Frame):
    """A step panel: number badge | title + description + status | action area."""

    def __init__(self, parent, number: int, title: str, desc: str):
        super().__init__(parent, bg=CARD, highlightbackground=BORDER,
                         highlightcolor=BORDER, highlightthickness=1)
        self.badge = tk.Label(self, text=str(number), font=FBOLD, fg="#FFFFFF",
                              bg="#C6CCD3", width=3, pady=4)
        self.badge.grid(row=0, column=0, rowspan=3, padx=(16, 14), pady=16, sticky="n")

        tk.Label(self, text=title, font=FBOLD, bg=CARD, fg=INK, anchor="w"
                 ).grid(row=0, column=1, sticky="w", pady=(16, 0))
        tk.Label(self, text=desc, font=FSUB, bg=CARD, fg=MUTE, anchor="w",
                 justify="left", wraplength=520
                 ).grid(row=1, column=1, sticky="w")
        self.status = tk.Label(self, text="", font=FSMALL, bg=CARD, fg=MUTE, anchor="w")
        self.status.grid(row=2, column=1, sticky="w", pady=(4, 16))

        self.actions = tk.Frame(self, bg=CARD)
        self.actions.grid(row=0, column=2, rowspan=3, padx=16, sticky="e")

        self.grid_columnconfigure(1, weight=1)

    def set_state(self, state: str, text: str = ""):
        colors = {
            "locked": ("#C6CCD3", MUTE),
            "ready":  (ACCENT, INK),
            "busy":   (WARN_BD, WARN_INK),
            "done":   (OK, OK),
            "error":  (BAD, BAD),
        }
        badge_bg, txt_fg = colors.get(state, ("#C6CCD3", MUTE))
        icon = {"done": "✓  ", "error": "✗  ", "busy": "…  "}.get(state, "")
        self.badge.configure(bg=badge_bg)
        if text:
            self.status.configure(text=icon + text, fg=txt_fg)


class App:
    def __init__(self, root: tk.Tk):
        self.root = root
        root.title("STRAFE Fee Calculator")
        root.geometry("860x720")
        root.minsize(780, 640)
        root.configure(bg=BG)

        style = ttk.Style(root)
        try:
            style.theme_use("clam")
        except tk.TclError:
            pass
        style.configure("M.Horizontal.TProgressbar", troughcolor="#E1E5E9",
                        background=ACCENT, bordercolor="#E1E5E9",
                        lightcolor=ACCENT, darkcolor=ACCENT, thickness=12)

        self.q: "queue.Queue" = queue.Queue()
        self.stop_event = threading.Event()
        self.cp_reply: "queue.Queue" = queue.Queue()
        self.done_reply: "queue.Queue" = queue.Queue()
        self.worker: threading.Thread | None = None
        self.xlsx_path: str | None = None

        # ---- header ------------------------------------------------
        hdr = tk.Frame(root, bg=HDR_BG)
        hdr.pack(fill="x")
        tk.Label(hdr, text="STRAFE Fee Calculator", font=FH1, bg=HDR_BG,
                 fg="#FFFFFF").pack(anchor="w", padx=22, pady=(16, 0))
        tk.Label(hdr, text="Fill the Testing Calculator from the test workbook, "
                 "then write the fees back.", font=FSUB, bg=HDR_BG, fg=HDR_SUB
                 ).pack(anchor="w", padx=22, pady=(2, 16))

        body = tk.Frame(root, bg=BG)
        body.pack(fill="both", expand=True, padx=18, pady=14)

        # ---- step 1 ----------------------------------------------
        self.c1 = Card(body, 1, "Connect Edge",
                       "Opens Microsoft Edge on the Testing Calculator with the "
                       "debugging port on. Log in to STRAFE in that window the "
                       "first time.")
        self.c1.pack(fill="x", pady=(0, 10))
        self.btn_edge = make_button(self.c1.actions, "Connect Edge", self.on_connect_edge, width=12)
        self.btn_edge.pack()

        # ---- step 2 ----------------------------------------------
        self.c2 = Card(body, 2, "Choose Excel file",
                       "Pick the .xlsx test workbook. It must have a STRAFE sheet "
                       "with data rows.")
        self.c2.pack(fill="x", pady=10)
        self.btn_file = make_button(self.c2.actions, "Choose file…", self.on_choose_file, width=12)
        self.btn_file.pack()

        # ---- step 3 ----------------------------------------------
        self.c3 = Card(body, 3, "Run it",
                       "Fills the calculator row by row. Pauses once at ~10% so "
                       "you can check. Stop any time – rerun continues where "
                       "it left off.")
        self.c3.pack(fill="x", pady=10)
        self.btn_run = make_button(self.c3.actions, "Run it", self.on_run, kind="go", width=12)
        self.btn_run.pack(pady=(0, 6))
        self.btn_stop = make_button(self.c3.actions, "Stop", self.on_stop, kind="danger", width=12)
        self.btn_stop.pack()
        self.chk_overwrite = tk.BooleanVar(value=False)
        tk.Checkbutton(self.c3, text="Redo everything (ignore saved results)",
                       variable=self.chk_overwrite, bg=CARD, fg=MUTE, font=FSMALL,
                       activebackground=CARD, selectcolor=CARD, bd=0,
                       highlightthickness=0, anchor="w"
                       ).grid(row=3, column=1, sticky="w", padx=0, pady=(0, 12))

        # ---- checkpoint bar (hidden until a checkpoint) ---------
        self.cp = tk.Frame(body, bg=WARN_BG, highlightbackground=WARN_BD,
                           highlightthickness=1)
        tk.Label(self.cp, text="Checkpoint – the file is open in Excel. "
                 "Review it, close Excel, then continue.", bg=WARN_BG, fg=WARN_INK,
                 font=FBOLD).pack(side="left", padx=14, pady=10)
        make_button(self.cp, "Continue", lambda: self._cp_answer("continue"),
                    kind="go").pack(side="left", padx=(0, 6), pady=8)
        make_button(self.cp, "Stop & review", lambda: self._cp_answer("stop"),
                    kind="ghost").pack(side="left", padx=(0, 12), pady=8)

        # ---- "are you satisfied?" bar (hidden until a run finishes) ----
        self.done_bar = tk.Frame(body, bg="#E6F4EA", highlightbackground=OK,
                                 highlightthickness=1)
        make_button(self.done_bar, "No – keep it",
                    lambda: self._done_answer(False), kind="ghost", width=12
                    ).pack(side="right", padx=(0, 12), pady=10)
        make_button(self.done_bar, "Yes – all done",
                    lambda: self._done_answer(True), kind="go", width=12
                    ).pack(side="right", padx=(0, 6), pady=10)
        self.done_msg = tk.Label(self.done_bar, text="", bg="#E6F4EA", fg="#0F5F2D",
                                 font=FBOLD, justify="left", anchor="w")
        self.done_msg.pack(side="left", fill="x", expand=True, padx=14, pady=10)

        # ---- progress -------------------------------------------
        prog = tk.Frame(body, bg=BG)
        prog.pack(fill="x", pady=(6, 6))
        tk.Label(prog, text="Progress", font=FSMALL, bg=BG, fg=MUTE).pack(side="left")
        self.pbar = ttk.Progressbar(prog, mode="determinate",
                                    style="M.Horizontal.TProgressbar")
        self.pbar.pack(side="left", fill="x", expand=True, padx=10)
        self.lbl_count = tk.Label(prog, text="–", font=FSMALL, bg=BG, fg=INK,
                                  width=12, anchor="e")
        self.lbl_count.pack(side="left")

        # ---- log ------------------------------------------------
        logwrap = tk.Frame(body, bg=CARD, highlightbackground=BORDER,
                           highlightthickness=1)
        logwrap.pack(fill="both", expand=True)
        self.log = ScrolledText(logwrap, height=9, wrap="word", state="disabled",
                                font=FMONO, bg=LOG_BG, fg=LOG_INK, bd=0,
                                relief="flat", padx=10, pady=8,
                                highlightthickness=0)
        self.log.pack(fill="both", expand=True)

        tk.Label(body, text="The workbook must be closed in Excel while a run is "
                 "in progress.", font=FSMALL, bg=BG, fg=MUTE).pack(anchor="w", pady=(8, 0))

        # initial states
        self.c1.set_state("ready", "Ready.")
        self._lock_step(self.c2, self.btn_file, "Do step 1 first.")
        self._lock_step(self.c3, self.btn_run, "Do step 2 first.")
        set_button(self.btn_stop, False)

        self.root.after(120, self._drain)

    # ---------------------------------------------------------------- ui utils

    def _lock_step(self, card: Card, *buttons):
        card.set_state("locked", buttons[-1] if isinstance(buttons[-1], str) else "")
        for b in buttons:
            if not isinstance(b, str):
                set_button(b, False)

    def _say(self, msg: str) -> None:
        self.log.configure(state="normal")
        self.log.insert("end", msg + "\n")
        self.log.see("end")
        self.log.configure(state="disabled")

    def _cp_answer(self, ans: str) -> None:
        self.cp.pack_forget()
        self.cp_reply.put(ans)

    def _done_answer(self, satisfied: bool) -> None:
        self.done_bar.pack_forget()
        self.done_reply.put(satisfied)

    # ---------------------------------------------------------------- step 1

    def on_connect_edge(self) -> None:
        set_button(self.btn_edge, False)
        self.c1.set_state("busy", "Opening Edge…")
        self._say("Step 1: launching Edge with the debugging port…")

        def work():
            try:
                how = fc.launch_debug_edge()
            except fc.SetupError as e:
                self.q.put(("edge_fail", {"msg": str(e)}))
                return
            self.q.put(("log", {"msg": "  (used launch_edge_debug.bat)"
                                if how == "bat" else "  (launched Edge directly)"}))
            ok = fc.verify_cdp(fc.CDP_URL_DEFAULT, timeout=30)
            self.q.put(("edge_ok" if ok else "edge_fail",
                        {"msg": "" if ok else
                         "Edge opened but the debugging port did not respond. "
                         "Close ALL Edge windows and try again."}))

        threading.Thread(target=work, daemon=True).start()

    # ---------------------------------------------------------------- step 2

    def on_choose_file(self) -> None:
        path = filedialog.askopenfilename(
            title="Choose the test workbook",
            filetypes=[("Excel workbook", "*.xlsx")],
        )
        if not path:
            return
        set_button(self.btn_file, False)
        self.c2.set_state("busy", "Checking…")
        self._say(f"Step 2: checking {os.path.basename(path)} …")

        def work():
            try:
                from openpyxl import load_workbook
                wb = load_workbook(path, read_only=True)
                if fc.SHEET_NAME_DEFAULT not in wb.sheetnames:
                    raise fc.SetupError(
                        f"the workbook has no '{fc.SHEET_NAME_DEFAULT}' sheet "
                        f"(sheets: {wb.sheetnames})")
                ws = wb[fc.SHEET_NAME_DEFAULT]
                n = len(fc.data_row_numbers(ws, fc.FIRST_DATA_ROW))
                wb.close()
                if n == 0:
                    raise fc.SetupError("no data rows found from row 23 down "
                                        "(column B is empty).")
                self.q.put(("file_ok", {"path": path, "rows": n}))
            except Exception as e:  # noqa: BLE001
                msg = str(e) if isinstance(e, fc.SetupError) else fc.explain_error(e)
                self.q.put(("file_fail", {"msg": msg}))

        threading.Thread(target=work, daemon=True).start()

    # ---------------------------------------------------------------- step 3

    def on_run(self) -> None:
        if not self.xlsx_path:
            return
        self.stop_event.clear()
        set_button(self.btn_run, False)
        set_button(self.btn_file, False)
        set_button(self.btn_edge, False)
        set_button(self.btn_stop, True)
        self.c3.set_state("busy", "Running…")
        self.pbar.configure(value=0)
        self.lbl_count.configure(text="–")
        self._say("\n" + "─" * 56 + "\nStep 3: run starting…\n")

        cfg = fc.EngineConfig(xlsx=self.xlsx_path,
                              overwrite=self.chk_overwrite.get(), checkpoint=True)
        hooks = GuiHooks(self.q, self.stop_event, self.cp_reply, self.done_reply)

        def work():
            try:
                fc.run_engine(cfg, hooks)
            except fc.SetupError as e:
                self.q.put(("fatal", {"msg": str(e)}))
            except ModuleNotFoundError:
                self.q.put(("fatal", {"msg": "Playwright is not installed. Run:\n"
                                      "  .venv\\Scripts\\python -m pip install playwright openpyxl"}))
            except Exception as e:  # noqa: BLE001
                self.q.put(("fatal", {"msg": fc.explain_error(e)}))
            finally:
                self.q.put(("worker_done", {}))

        self.worker = threading.Thread(target=work, daemon=True)
        self.worker.start()

    def on_stop(self) -> None:
        self.stop_event.set()
        set_button(self.btn_stop, False)
        self.c3.set_state("busy", "Stopping after this row…")
        self._say("Stop requested – will stop after the current row…")

    def _finish_ui(self, state="done", text="Finished.") -> None:
        set_button(self.btn_stop, False)
        set_button(self.btn_run, True)
        self.btn_run.configure(text="Run again")
        set_button(self.btn_file, True)
        set_button(self.btn_edge, True)
        self.cp.pack_forget()
        self.c3.set_state(state, text)

    # ---------------------------------------------------------------- queue pump

    def _drain(self) -> None:
        try:
            while True:
                kind, d = self.q.get_nowait()
                self._handle(kind, d)
        except queue.Empty:
            pass
        self.root.after(120, self._drain)

    def _handle(self, kind: str, d: dict) -> None:
        if kind == "log":
            self._say(d["msg"])

        elif kind == "edge_ok":
            self.c1.set_state("done", "Connected. Keep the Testing Calculator open "
                                      "(log in if asked).")
            self._say("Step 1: OK – Edge is reachable.\n")
            set_button(self.btn_edge, True)
            self.btn_edge.configure(text="Reconnect")
            set_button(self.btn_file, True)
            self.c2.set_state("ready", "Ready.")
        elif kind == "edge_fail":
            self.c1.set_state("error", d["msg"] or "Could not connect.")
            self._say(f"Step 1 FAILED: {d['msg']}\n")
            set_button(self.btn_edge, True)

        elif kind == "file_ok":
            self.xlsx_path = d["path"]
            self.c2.set_state("done", f"{os.path.basename(d['path'])}  ·  {d['rows']} rows")
            self._say(f"Step 2: OK – {d['rows']} data rows.\n")
            set_button(self.btn_file, True)
            self.btn_file.configure(text="Change file…")
            set_button(self.btn_run, True)
            self.c3.set_state("ready", "Ready to run.")
        elif kind == "file_fail":
            self.c2.set_state("error", d["msg"])
            self._say(f"Step 2 FAILED: {d['msg']}\n")
            set_button(self.btn_file, True)

        elif kind == "start":
            total = max(d["to_process"], 1)
            self.pbar.configure(maximum=total, value=0)
            self.lbl_count.configure(text=f"0 / {d['to_process']}")
        elif kind == "progress":
            self.pbar.configure(value=d["done"])
            self.lbl_count.configure(text=f"{d['done']} / {d['total']}")

        elif kind == "row_begin":
            self._say(f"row {d['row']}: {d['quick_ref']}   {d['liability']} | {d['coverage']}")
        elif kind == "pass_done":
            tag = "    (message, no fee)" if d["is_message"] else ""
            self._say(f"      {d['col']}  →  {d['kcol']}{d['row']} = {d['value']}{tag}")
        elif kind == "row_skip":
            self._say(f"row {d['row']}: {d['reason']}")
        elif kind == "row_done_already":
            self._say(f"row {d['row']}: already filled – skipping")
        elif kind == "row_changed":
            self._say(f"row {d['row']}: inputs changed – redoing this row")
        elif kind == "row_failed":
            self._say(f"row {d['row']}: FAILED while {d['stage']} – {d['friendly']}")
        elif kind == "saved":
            pass

        elif kind == "checkpoint":
            self._say("\n" + "─" * 56)
            self._say(f"CHECKPOINT: {d['done']} of {d['total']} rows done and saved.")
            self._say("Opening the file in Excel – check it, close Excel, then Continue.")
            self._say("─" * 56)
            self.c3.set_state("busy", "Checkpoint – waiting for you.")
            try:
                os.startfile(d["path"])  # noqa: S606
            except Exception as e:  # noqa: BLE001
                self._say(f"(could not open Excel automatically: {e})")
            self.cp.pack(fill="x", pady=(4, 8), before=self.pbar.master)
        elif kind == "resumed":
            self.cp.pack_forget()
            self.c3.set_state("busy", "Running…")
            self._say("Resuming…\n")

        elif kind == "stopped":
            self._say("\nSTOPPED. Progress saved. Edit the sheet if you like, then "
                      "click 'Run again' – it continues from here and redoes "
                      "only rows you changed.")
            self._finish_ui("ready", "Stopped – click Run again to continue.")
        elif kind == "aborted":
            self._say("\n" + "═" * 56)
            self._say(f"STOPPED at row {d['row']} while {d['stage']}.")
            self._say(f"Reason: {d['friendly']}")
            self._say("Progress saved. Fix the cause, then click 'Run again'.")
            self._say("═" * 56)
            self._finish_ui("error", f"Stopped at row {d['row']}.")
        elif kind == "finished":
            s = d["summary"]
            self._say(f"\nDone: {s['processed']} processed, {s['already_done']} "
                      f"already filled, {s['skipped']} skipped, "
                      f"{len(s['failed_rows'])} failed, {s['messages']} message(s).")
            if s["failed_rows"]:
                self._say(f"Failed rows (Run again to retry): {s['failed_rows']}")
        elif kind == "ask_satisfied":
            s = d["summary"]
            if s.get("messages"):
                self._say("Note: some rows show a message instead of a fee – "
                          "check those before saying you're done.")
            self.done_msg.configure(
                text="Results look good?\n"
                     "Yes = delete the resume file    No = keep it, edit & Run again")
            self.done_bar.pack(fill="x", pady=(4, 8), before=self.pbar.master)
            self.c3.set_state("busy", "Waiting for your answer…")
        elif kind == "state_deleted":
            self._say("Resume file deleted – this run is marked done.")
        elif kind == "state_kept":
            self._say("Resume file kept – edit rows and Run again to continue.")

        elif kind == "fatal":
            self._say(f"\nERROR: {d['msg']}")
            self._finish_ui("error", "Error – see the log.")
        elif kind == "worker_done":
            # normal finish (stopped/aborted already reset the UI)
            if str(self.btn_run["state"]) == "disabled":
                self._finish_ui("done", "Finished.")


class GuiHooks(fc.Hooks):
    def __init__(self, q: "queue.Queue", stop_event: threading.Event,
                 cp_reply: "queue.Queue", done_reply: "queue.Queue"):
        self.q = q
        self.stop_event = stop_event
        self.cp_reply = cp_reply
        self.done_reply = done_reply

    def log(self, msg: str) -> None:
        self.q.put(("log", {"msg": msg}))

    def event(self, kind: str, **data) -> None:
        self.q.put((kind, data))

    def should_stop(self) -> bool:
        return self.stop_event.is_set()

    def wait_at_checkpoint(self) -> str:
        return self.cp_reply.get()

    def ask_satisfied(self, summary: dict) -> bool:
        self.q.put(("ask_satisfied", {"summary": summary}))
        return self.done_reply.get()


def _enable_dpi_awareness() -> None:
    """Crisp text on high-DPI / scaled displays (Windows)."""
    try:
        from ctypes import windll
        try:
            windll.shcore.SetProcessDpiAwareness(2)   # per-monitor v1
        except Exception:
            windll.user32.SetProcessDPIAware()
    except Exception:
        pass


def main() -> None:
    _enable_dpi_awareness()
    root = tk.Tk()
    try:
        root.tk.call("tk", "scaling", root.winfo_fpixels("1i") / 72.0)
    except tk.TclError:
        pass
    App(root)
    root.mainloop()


def _main_guarded() -> None:
    """Entry point for the packaged .exe: a startup crash must not vanish
    silently in a windowed build - log it and show a dialog."""
    try:
        main()
    except Exception as e:  # noqa: BLE001
        import traceback
        log_path = os.path.join(fc.app_data_dir(), "last-error.log")
        try:
            with open(log_path, "w", encoding="utf-8") as f:
                f.write(traceback.format_exc())
        except OSError:
            pass
        try:
            from tkinter import messagebox
            messagebox.showerror("STRAFE Fee Calculator - startup error",
                                 f"{e}\n\nDetails saved to:\n{log_path}")
        except Exception:
            print(traceback.format_exc(), flush=True)
        raise


if __name__ == "__main__":
    _main_guarded()