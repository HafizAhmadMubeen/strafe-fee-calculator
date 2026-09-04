r"""
fill_calculator.py
==================================================================================
Drive the STRAFE Testing Calculator from the "STRAFE" worksheet of the endorsement
test workbook, then write the resulting Fee Amount back into the sheet.

Use the GUI (fill_calculator_gui.py) for a click-through version, or this script
directly from the command line.

WHAT IT DOES  (per data row, row 23 down while column B "Quick Ref" has a value)
--------------------------------------------------------------------------------
  Globals (set once, from column E at the top of the sheet):
    E2 -> State        E5 -> Underwriter        E8 -> Closing Date (as YYYY-MM-DD)

  Per row:
    B  -> Fee Item Code
    A  -> "Select Liability Type" dropdown (row 1 of the coverage group)
    AD -> "Coverage Type Code"            (row 1 of the coverage group)
    V  -> "Boolean Question" (row 1)          - only if V has a value
    W  -> "Select Answer" dropdown (row 1)    - only if V has a value  (Yes / No)

  For each of G / H / I / J that has a value in the row, type the ROW-22 header
  amount for that column into "Amount" (G->1000  H->4000000  I->30000000
  J->51000000), click Calculate, and write the Fee Amount as a NUMBER into the
  matching Actual Value column (G->K  H->L  I->M  J->N).  If the calculator
  returns no fee, the Response/Error Msg text is written there instead.

  A row is skipped (reason written into column K) when A or AD is blank, A is not
  a valid dropdown option, or none of G/H/I/J has a value.

RESUME / CHANGE DETECTION
--------------------------------------------------------------------------------
  A sidecar "<name> (state).json" records each row's input fingerprint.  Rerunning
  the same command:
    - skips rows whose Actual Value cells are already filled AND whose inputs are
      unchanged,
    - RE-RUNS a row whose inputs you edited (even if its result was filled),
    - fills any result cell you cleared.
  So after a Stop you can edit the sheet and just run again.

CHECKPOINT
--------------------------------------------------------------------------------
  After the first ~10% of the rows it will process (only if that is >= 5 rows),
  it saves and pauses so you can eyeball the results.  CLI: press Enter to go on,
  or type s + Enter to stop.

BROWSER
--------------------------------------------------------------------------------
  The script attaches to an Edge tab you already have open on the Testing
  Calculator (Edge started with --remote-debugging-port=9222).  It never opens,
  closes, reloads or navigates the browser.  Use launch_edge_debug.bat, or the
  GUI's "Connect Edge" button, to start that Edge.

SETUP (one time)
--------------------------------------------------------------------------------
    python -m venv .venv
    .venv\Scripts\python -m pip install playwright openpyxl

RUN
--------------------------------------------------------------------------------
    .venv\Scripts\python fill_calculator.py "<workbook>.xlsx" [options]

    --limit N        process at most N rows this run
    --start-row R    first sheet row (default 23)
    --sheet NAME     worksheet name (default STRAFE)
    --overwrite      redo every row, ignore existing values / state
    --retry-messages redo cells that hold a text message instead of a $ amount
    --keep-going     on a row error, log it and continue instead of stopping
    --no-checkpoint  do not pause at ~10%
    --no-backup      do not make the "<name> (backup).xlsx" safety copy
    --cdp-url URL    CDP endpoint (default http://127.0.0.1:9222)
    --settle-ms MS   wait budget for the result after Calculate (default 4000)
    --dry-run        print the per-row plan; do NOT touch the browser
    --list-fields    dump every field on the page with its id, then exit
"""

from __future__ import annotations

import argparse
import datetime as dt
import hashlib
import json
import math
import os
import shutil
import subprocess
import sys
import time
import urllib.request

# ------------------------------------------------------------------ config ----

SHEET_NAME_DEFAULT = "STRAFE"
CDP_URL_DEFAULT = "http://127.0.0.1:9222"   # NOT "localhost" - can resolve to ::1
PAGE_URL_MUST_CONTAIN = "strafe.stewart.com/TestingCalculator"
CALCULATOR_URL = "https://strafe.stewart.com/TestingCalculator"

HEADER_ROW = 21
AMOUNT_ROW = 22          # holds the dollar amount for each G/H/I/J column
FIRST_DATA_ROW = 23

CELL_STATE = "E2"
CELL_UNDERWRITER = "E5"
CELL_EFFECTIVE_DATE = "E8"

COL_LIABILITY = "A"
COL_QUICK_REF = "B"
COL_QUESTION = "V"
COL_ANSWER = "W"
COL_COVERAGE_TYPE = "AD"

# Expected-Value column -> (column holding the Amount in AMOUNT_ROW,
#                           Actual-Value column to write the Fee Amount into)
PASS_COLUMNS = {
    "G": ("G", "K"),
    "H": ("H", "L"),
    "I": ("I", "M"),
    "J": ("J", "N"),
}

VALID_LIABILITY_TYPES = {
    "Binder Amount", "Commitment Amount", "Construction Loan Amount",
    "Construction Premium So Far", "Guarantee Amount", "Leasehold Amount",
    "Leasehold Loan Amount", "LMP Premium So Far", "Loan Amount",
    "LOP Premium So Far", "Mortgage Premium So Far", "Owner's Premium So Far",
    "Project Loan Amount", "Project Premium So Far", "Sales Price",
    "True/Full Rate CMP Premium", "True/Full Rate LMP Premium",
    "True/Full Rate Mortgage Premium", "True/Full Rate Owners Premium",
    "True/Full Rate PMP Premium", "Unpaid Principal Balance",
}

CHECKPOINT_FRACTION = 0.10
CHECKPOINT_MIN_ROWS = 5
SAVE_EVERY_ROWS = 10

def base_dir() -> str:
    """Folder the app lives in: the .exe's folder when frozen, else this file's."""
    if getattr(sys, "frozen", False):
        return os.path.dirname(sys.executable)
    return os.path.dirname(os.path.abspath(__file__))


def app_data_dir() -> str:
    """Per-user folder for our own files (always user-writable, unlike a
    Program Files exe dir).  Used for the error log and the fallback Edge profile."""
    base = os.environ.get("LOCALAPPDATA") or os.path.expanduser("~")
    d = os.path.join(base, "STRAFE Fee Calculator")
    try:
        os.makedirs(d, exist_ok=True)
    except OSError:
        d = os.path.expanduser("~")
    return d


# Path to the IT-approved launcher, if it sits next to us.
EDGE_BAT = os.path.join(base_dir(), "launch_edge_debug.bat")

# Dedicated Edge profile so we never disturb the user's normal Edge, and the
# STRAFE login is remembered between runs.  When launch_edge_debug.bat is used it
# sets this same '.edge-automation' folder itself, so there is ONE login.
if getattr(sys, "frozen", False):
    EDGE_PROFILE_DIR = os.path.join(app_data_dir(), "edge-profile")
else:
    EDGE_PROFILE_DIR = os.path.join(base_dir(), ".edge-automation")

# ---------------------------------------------------------------- selectors ----
# Pinned to the live DOM ids (confirmed via --list-fields).  If the site changes,
# run --list-fields and fix the one list that is wrong.

SELECTORS = {
    "closing_date": ['css:#ClosingDate'],
    "underwriter": ['css:#Underwriter'],
    "state": ['css:#PropertyInformation_PropertyState'],          # not the SalesTax one
    "fee_item_code": ['css:#feeItemCodeInput'],
    "liability_type": ['css:[id="BasisInfo[0].IdAmountName"]'],    # <select>
    "coverage_type_code": ['css:[id="BasisInfo[0].BaseTitleInsTypeCode"]'],
    "amount": ['css:[id="BasisInfo[0].ValueStr"]'],
    "boolean_question": ['css:[id="QuestionList[0].Code"]'],
    "boolean_answer": ['css:[id="QuestionList[0].AnswerType"]'],   # <select> Select Answer/Yes/No
    "calculate_button": ['css:#btnTestToolCalculate'],
    "fee_amount": ['css:#txtFeeAmount'],
    "error_msg": ['css:#taResponseErrorMsg'],
}

LOADER_SELECTOR = "#loader-wrapper"


class SetupError(Exception):
    """A problem the user must fix before the run can start (bad file, no Edge…)."""


# ------------------------------------------------------------------ helpers ----

def is_blank(v) -> bool:
    return v is None or (isinstance(v, str) and v.strip() == "")


def as_text(v) -> str:
    return "" if v is None else str(v).strip()


def parse_money(s):
    """'$1,234.50' -> 1234.5 ; '' / junk -> None."""
    s = as_text(s).replace("$", "").replace(",", "").replace("%", "").strip()
    if s == "":
        return None
    try:
        return float(s)
    except ValueError:
        return None


def clean_amount(v) -> str:
    """Render a numeric cell as a plain string: 1000.0 -> '1000'."""
    if isinstance(v, bool):
        return str(v)
    if isinstance(v, int):
        return str(v)
    if isinstance(v, float):
        return str(int(v)) if v.is_integer() else repr(v)
    return as_text(v).replace(",", "").lstrip("$").strip()


def _to_date(v):
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    if isinstance(v, (int, float)):
        return (dt.datetime(1899, 12, 30) + dt.timedelta(days=float(v))).date()
    s = as_text(v)
    if s.isdigit():
        return (dt.datetime(1899, 12, 30) + dt.timedelta(days=int(s))).date()
    for fmt in ("%m/%d/%Y", "%m/%d/%y", "%Y-%m-%d"):
        try:
            return dt.datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    return None


def excel_date_to_iso(v) -> str:
    """The Closing Date field is <input type='date'> which needs YYYY-MM-DD."""
    d = _to_date(v)
    return d.strftime("%Y-%m-%d") if d else as_text(v)


def explain_error(e: Exception) -> str:
    """Turn a raw exception into one plain sentence about what went wrong."""
    name = type(e).__name__
    msg = str(e)
    low = msg.lower()
    if "loader-wrapper" in low or (name == "TimeoutError" and "click" in low):
        return ("the site's loading spinner never cleared - the calculator was too "
                "slow or got stuck. Run again to resume; if it repeats, raise the "
                "settle time or reload the Testing Calculator page.")
    if name == "TimeoutError":
        return ("timed out waiting for the page to respond - the site may be slow "
                "or a field did not appear. Run again to resume.")
    if isinstance(e, LookupError):
        return (f"{msg} - a field could not be found on the page. The site layout "
                f"may have changed; run --list-fields and update SELECTORS.")
    if ("target closed" in low or "websocket" in low or "browser has been closed" in low
            or "connection closed" in low):
        return ("lost the connection to Edge - the tab was closed or navigated away. "
                "Reopen the Testing Calculator in the debug Edge window and run again.")
    if isinstance(e, PermissionError):
        return "could not write the workbook - it is open in Excel. Close it and run again."
    return f"{name}: {msg}"


# --------------------------------------------------------------- state file ----

def state_path_for(xlsx: str) -> str:
    root, _ext = os.path.splitext(xlsx)
    return f"{root} (state).json"


def load_state(xlsx: str) -> dict:
    p = state_path_for(xlsx)
    try:
        with open(p, "r", encoding="utf-8") as f:
            data = json.load(f)
        if isinstance(data, dict) and isinstance(data.get("rows"), dict):
            return data
    except (OSError, ValueError):
        pass
    return {"version": 1, "rows": {}}


def save_state(xlsx: str, state: dict) -> None:
    p = state_path_for(xlsx)
    tmp = p + ".tmp"
    try:
        with open(tmp, "w", encoding="utf-8") as f:
            json.dump(state, f, indent=1)
        os.replace(tmp, p)
    except OSError:
        pass


def globals_signature(closing_date: str, underwriter: str, state_val: str) -> str:
    return f"{closing_date}|{underwriter}|{state_val}"


def row_signature(inp: dict, gsig: str) -> str:
    payload = {
        "g": gsig,
        "liability": inp["liability"],
        "coverage": inp["coverage"],
        "question": inp["question"],
        "answer": inp["answer"],
        "expected": {k: as_text(inp["expected"][k]) for k in PASS_COLUMNS},
    }
    blob = json.dumps(payload, sort_keys=True).encode("utf-8")
    return hashlib.sha1(blob).hexdigest()[:16]


# ---------------------------------------------------------------- backups ------

def backup_path_for(path: str) -> str:
    root, ext = os.path.splitext(path)
    return f"{root} (backup){ext}"


def make_backup(path: str):
    """Copy to '<name> (backup).xlsx'.  Deleted on a clean finish; kept on failure.
    Returns (backup_path, is_fresh)."""
    dst = backup_path_for(path)
    if os.path.exists(dst):
        return dst, False
    shutil.copy2(path, dst)
    return dst, True


def assert_not_locked(path: str) -> None:
    folder, name = os.path.split(os.path.abspath(path))
    if os.path.exists(os.path.join(folder, "~$" + name)):
        raise SetupError(f"'{name}' is open in Excel (lock file ~${name}). Close it and try again.")
    try:
        with open(path, "rb+"):
            pass
    except PermissionError:
        raise SetupError(f"cannot write '{path}' - close it in Excel and try again.")
    except OSError as e:
        raise SetupError(f"cannot open '{path}': {e}")


# --------------------------------------------------------------- edge launch ---

EDGE_CANDIDATES = [
    os.path.expandvars(r"%ProgramFiles(x86)%\Microsoft\Edge\Application\msedge.exe"),
    os.path.expandvars(r"%ProgramFiles%\Microsoft\Edge\Application\msedge.exe"),
    os.path.expandvars(r"%LocalAppData%\Microsoft\Edge\Application\msedge.exe"),
]


def find_edge() -> str | None:
    for p in EDGE_CANDIDATES:
        if p and os.path.isfile(p):
            return p
    return None


def launch_debug_edge(url: str = CALCULATOR_URL, profile_dir: str = EDGE_PROFILE_DIR,
                      port: int = 9222) -> str:
    """Start Edge with the debugging port.  Prefers launch_edge_debug.bat when it
    is next to us (it is already approved by IT and sets the same profile folder);
    otherwise launches msedge.exe directly.  Returns which path was used."""
    no_window = getattr(subprocess, "CREATE_NO_WINDOW", 0)

    if os.path.isfile(EDGE_BAT):
        subprocess.Popen(["cmd", "/c", EDGE_BAT], cwd=os.path.dirname(EDGE_BAT),
                         creationflags=no_window, close_fds=True)
        return "bat"

    exe = find_edge()
    if not exe:
        raise SetupError("could not find msedge.exe - is Microsoft Edge installed?")
    os.makedirs(profile_dir, exist_ok=True)
    subprocess.Popen(
        [exe, f"--remote-debugging-port={port}", f"--user-data-dir={profile_dir}", url],
        close_fds=True,
    )
    return "direct"


def verify_cdp(cdp_url: str = CDP_URL_DEFAULT, timeout: float = 25.0) -> bool:
    base = cdp_url.rstrip("/")
    urls = {base,
            base.replace("localhost", "127.0.0.1"),
            base.replace("127.0.0.1", "localhost")}
    deadline = time.time() + timeout
    while time.time() < deadline:
        for u in urls:
            try:
                with urllib.request.urlopen(u + "/json/version", timeout=2) as r:
                    if r.status == 200:
                        return True
            except Exception:
                pass
        time.sleep(1.0)
    return False


# ---------------------------------------------------------- playwright glue ----

def _locate(page, spec):
    if spec.startswith("label:"):
        return page.get_by_label(spec[len("label:"):], exact=False)
    if spec.startswith("css:"):
        return page.locator(spec[len("css:"):])
    if spec.startswith("xpath="):
        return page.locator(spec)
    if spec.startswith("xpath:"):
        return page.locator("xpath=" + spec[len("xpath:"):])
    if spec.startswith("option:"):
        return page.locator("select").filter(has_text=spec[len("option:"):])
    if spec.startswith("role:"):
        _, role, name = spec.split(":", 2)
        return page.get_by_role(role, name=name, exact=False)
    return page.locator(spec)


def field(page, key):
    last_err = None
    for spec in SELECTORS[key]:
        try:
            loc = _locate(page, spec)
            if loc.count() > 0:
                return loc.first
        except Exception as e:  # noqa: BLE001
            last_err = e
    raise LookupError(
        f"Could not locate field '{key}'. Tried: {SELECTORS[key]}. "
        f"Run with --list-fields and adjust SELECTORS['{key}']. ({last_err})"
    )


def wait_not_loading(page, timeout_ms: int = 60000) -> None:
    """Block while the site's full-page loading overlay is up. 'hidden' is also
    satisfied when the element is absent, so this is a no-op if there is no loader."""
    try:
        page.locator(LOADER_SELECTOR).first.wait_for(state="hidden", timeout=timeout_ms)
    except Exception:
        pass


def set_text(page, key, value):
    wait_not_loading(page)
    loc = field(page, key)
    loc.scroll_into_view_if_needed()
    value = "" if value is None else str(value)
    try:
        input_type = (loc.get_attribute("type") or "").lower()
    except Exception:
        input_type = ""
    if input_type == "date" and value:
        value = excel_date_to_iso(value)
    loc.click()
    loc.fill("")
    if value != "":
        loc.fill(value)
    loc.press("Tab")


def set_select(page, key, value):
    wait_not_loading(page)
    loc = field(page, key)
    loc.scroll_into_view_if_needed()
    try:
        loc.select_option(label=value)
    except Exception:
        loc.select_option(value)
    loc.press("Tab")


def read_value(loc) -> str:
    try:
        v = loc.input_value()
        if v is not None and v.strip() != "":
            return v.strip()
    except Exception:
        pass
    try:
        return (loc.inner_text() or "").strip()
    except Exception:
        return ""


def connect_page(cdp_url: str):
    from playwright.sync_api import sync_playwright

    candidates = [cdp_url]
    if "localhost" in cdp_url:
        candidates.append(cdp_url.replace("localhost", "127.0.0.1"))
    if "127.0.0.1" in cdp_url:
        candidates.append(cdp_url.replace("127.0.0.1", "localhost"))

    pw = sync_playwright().start()
    browser = None
    last_err = None
    for url in candidates:
        try:
            browser = pw.chromium.connect_over_cdp(url)
            break
        except Exception as e:  # noqa: BLE001
            last_err = e
    if browser is None:
        pw.stop()
        raise SetupError(
            "could not attach to Edge on the debugging port.\n"
            "  1. Close every Edge window.\n"
            "  2. Use 'Connect Edge' (GUI) or launch_edge_debug.bat.\n"
            "  3. Open http://127.0.0.1:9222/json/version - you must see JSON.\n"
            f"  ({last_err})"
        )

    target = None
    for ctx in browser.contexts:
        for pg in ctx.pages:
            if PAGE_URL_MUST_CONTAIN.lower() in (pg.url or "").lower():
                target = pg
                break
        if target:
            break
    if target is None:
        open_urls = [pg.url for ctx in browser.contexts for pg in ctx.pages]
        pw.stop()
        raise SetupError(
            f"no open tab is on the Testing Calculator ({PAGE_URL_MUST_CONTAIN}).\n"
            f"Open tabs: {open_urls}\n"
            "In the debug Edge window, log in and open the Testing Calculator, then retry."
        )

    target.bring_to_front()
    return pw, browser, target


def list_fields(page) -> None:
    js = r"""
    () => {
      const near = (el) => {
        let n = el, hops = 0;
        while (n && hops < 6) {
          let p = n.previousElementSibling;
          while (p) {
            const t = (p.innerText || p.textContent || '').trim();
            if (t) return t.slice(0, 60);
            p = p.previousElementSibling;
          }
          n = n.parentElement; hops++;
        }
        return '';
      };
      const out = [];
      document.querySelectorAll('input,select,textarea,button').forEach((el, i) => {
        out.push({
          i, tag: el.tagName.toLowerCase(), type: el.type || '',
          id: el.id || '', name: el.name || '', placeholder: el.placeholder || '',
          value: (el.value || '').slice(0, 40),
          options: el.tagName.toLowerCase() === 'select'
            ? [...el.options].map(o => o.text.trim()).slice(0, 6) : undefined,
          label: near(el),
        });
      });
      return out;
    }
    """
    rows = page.evaluate(js)
    print(f"--- {len(rows)} form controls on {page.url} ---", flush=True)
    for r in rows:
        extra = f"  options={r['options']}" if r.get("options") else ""
        print(f"[{r['i']:>3}] {r['tag']:<8} type={r['type']:<10} id={r['id']:<26} "
              f"name={r['name']:<26} label='{r['label']}' val='{r['value']}'{extra}",
              flush=True)


# ----------------------------------------------------------------- planning ----

def data_row_numbers(ws, start_row):
    r = start_row
    out = []
    while not is_blank(ws[f"{COL_QUICK_REF}{r}"].value):
        out.append(r)
        r += 1
    return out


def read_row_inputs(ws, r) -> dict:
    return {
        "quick_ref": as_text(ws[f"{COL_QUICK_REF}{r}"].value),
        "liability": as_text(ws[f"{COL_LIABILITY}{r}"].value),
        "coverage": as_text(ws[f"{COL_COVERAGE_TYPE}{r}"].value),
        "question": as_text(ws[f"{COL_QUESTION}{r}"].value),
        "answer": as_text(ws[f"{COL_ANSWER}{r}"].value),
        "expected": {col: ws[f"{col}{r}"].value for col in PASS_COLUMNS},
    }


def skip_reason(inp, amounts):
    reasons = []
    if is_blank(inp["liability"]):
        reasons.append("Liability Type missing")
    elif inp["liability"] not in VALID_LIABILITY_TYPES:
        reasons.append(f'Liability Type "{inp["liability"]}" not a valid option')
    if is_blank(inp["coverage"]):
        reasons.append("Coverage Type Code missing")
    if not amounts:
        reasons.append("No Expected Value")
    return "Skipped: " + "; ".join(reasons) if reasons else None


def read_globals(ws):
    state_val = as_text(ws[CELL_STATE].value)
    underwriter = as_text(ws[CELL_UNDERWRITER].value)
    closing_date = excel_date_to_iso(ws[CELL_EFFECTIVE_DATE].value)
    amount_for = {
        col: clean_amount(ws[f"{ac}{AMOUNT_ROW}"].value)
        for col, (ac, _kc) in PASS_COLUMNS.items()
    }
    return closing_date, underwriter, state_val, amount_for


# ------------------------------------------------------------------- hooks -----

class Hooks:
    """Sink for engine events.  Subclass for CLI / GUI."""

    def log(self, msg: str) -> None:
        print(msg, flush=True)

    def event(self, kind: str, **data) -> None:
        pass

    def should_stop(self) -> bool:
        return False

    def wait_at_checkpoint(self) -> str:
        """Return 'continue' or 'stop'."""
        return "continue"

    def ask_satisfied(self, summary: dict) -> bool:
        """Called once after a run that completed. True -> the resume/state file
        is deleted (user is done); False -> it is kept (user may edit + rerun).
        Default keeps it (safe)."""
        return False


class ConsoleHooks(Hooks):
    def event(self, kind, **d):
        if kind == "start":
            self.log(f"Rows: {d['total_rows']} found, {d['to_process']} to process "
                     f"this run.\n")
        elif kind == "row_begin":
            self.log(f"row {d['row']}: {d['quick_ref']}  {d['liability']} | {d['coverage']}")
        elif kind == "pass_done":
            tag = "  (message, no fee)" if d["is_message"] else ""
            self.log(f"    {d['col']} amount={d['amount']} -> {d['kcol']}{d['row']} "
                     f"= {d['value']!r}{tag}")
        elif kind == "row_skip":
            self.log(f"row {d['row']}: {d['reason']}")
        elif kind == "row_done_already":
            self.log(f"row {d['row']}: {d['quick_ref']}  already filled - skipping")
        elif kind == "row_changed":
            self.log(f"row {d['row']}: inputs changed - redoing")
        elif kind == "row_failed":
            self.log(f"row {d['row']}: FAILED while {d['stage']} - {d['friendly']}")
        elif kind == "saved":
            self.log(f"  (saved: {d['path']})")
        elif kind == "checkpoint":
            self.log("\n" + "-" * 68)
            self.log(f"CHECKPOINT - {d['done']} of {d['total']} rows done and saved.")
            self.log(f"Open and check:  {d['path']}")
            self.log("-" * 68)
        elif kind == "resumed":
            self.log("Resuming...\n")
        elif kind == "stopped":
            self.log("\nStopped. Progress saved - run the same command to continue.")
        elif kind == "aborted":
            self.log("\n" + "=" * 68)
            self.log(f"STOPPED at row {d['row']} while {d['stage']}.")
            self.log(f"Reason: {d['friendly']}")
            self.log("\nProgress saved. Fix the cause, then run the SAME command "
                     "again - it resumes from here.")
            self.log("=" * 68)
        elif kind == "finished":
            s = d["summary"]
            self.log(f"\nSummary: {s['processed']} processed, "
                     f"{s['already_done']} already filled, {s['skipped']} skipped, "
                     f"{len(s['failed_rows'])} failed, "
                     f"{s['messages']} message(s) (no fee).")
            if s["failed_rows"]:
                self.log(f"Failed rows (run again to retry): {s['failed_rows']}")
        elif kind == "state_deleted":
            self.log("Resume file deleted - this run is marked done.")
        elif kind == "state_kept":
            self.log("Resume file kept - edit rows and run again to continue.")

    def ask_satisfied(self, summary: dict) -> bool:
        if not sys.stdin or not sys.stdin.isatty():
            return False
        try:
            ans = input("\nResults look good? Delete the resume file so the folder "
                        "is clean? [y/N]: ").strip().lower()
        except EOFError:
            return False
        return ans in ("y", "yes")

    def wait_at_checkpoint(self) -> str:
        if not sys.stdin or not sys.stdin.isatty():
            self.log("(non-interactive - continuing automatically)")
            return "continue"
        try:
            ans = input("Press Enter to continue, or type s + Enter to stop: ").strip().lower()
        except EOFError:
            return "continue"
        return "stop" if ans == "s" else "continue"


# ------------------------------------------------------------------ config -----

class EngineConfig:
    def __init__(self, xlsx, sheet=SHEET_NAME_DEFAULT, start_row=FIRST_DATA_ROW,
                 limit=None, cdp_url=CDP_URL_DEFAULT, settle_ms=4000,
                 overwrite=False, retry_messages=False, keep_going=False,
                 checkpoint=True, no_backup=False):
        self.xlsx = xlsx
        self.sheet = sheet
        self.start_row = start_row
        self.limit = limit
        self.cdp_url = cdp_url
        self.settle_ms = settle_ms
        self.overwrite = overwrite
        self.retry_messages = retry_messages
        self.keep_going = keep_going
        self.checkpoint = checkpoint
        self.no_backup = no_backup


# ------------------------------------------------------------------ engine -----

def _pending_passes(ws, rownum, amounts, inputs_changed, cfg):
    out = []
    for (col, amt, kcol) in amounts:
        cv = ws[f"{kcol}{rownum}"].value
        if (cfg.overwrite or is_blank(cv) or inputs_changed
                or (cfg.retry_messages and not isinstance(cv, (int, float)))):
            out.append((col, amt, kcol))
    return out


def _filled_cols(ws, rownum, amounts):
    return [kc for (_c, _a, kc) in amounts
            if not is_blank(ws[f"{kc}{rownum}"].value)]


def run_engine(cfg: EngineConfig, hooks: Hooks) -> dict:
    """Do the run. Returns a summary dict. Setup problems raise SetupError."""
    from openpyxl import load_workbook

    xlsx = cfg.xlsx
    if not os.path.isfile(xlsx):
        raise SetupError(f"file not found: {xlsx}")

    wb = load_workbook(xlsx)
    if cfg.sheet not in wb.sheetnames:
        raise SetupError(f"no sheet '{cfg.sheet}'. Sheets: {wb.sheetnames}")
    ws = wb[cfg.sheet]

    state = load_state(xlsx)
    rows = data_row_numbers(ws, cfg.start_row)
    closing_date, underwriter, state_val, amount_for = read_globals(ws)
    gsig = globals_signature(closing_date, underwriter, state_val)

    hooks.log(f"Workbook : {xlsx}")
    hooks.log(f"Sheet    : {cfg.sheet}")
    hooks.log(f"Globals  : State={state_val!r} Underwriter={underwriter!r} "
              f"ClosingDate={closing_date!r}")

    # ---- count how many rows this run will actually process --------------
    to_process_rows = []
    for r in rows:
        inp = read_row_inputs(ws, r)
        amounts = [(c, amount_for[c], PASS_COLUMNS[c][1])
                   for c in PASS_COLUMNS if not is_blank(inp["expected"][c])]
        if skip_reason(inp, amounts):
            continue
        stored = state["rows"].get(str(r))
        changed = stored is not None and stored.get("sig") != row_signature(inp, gsig)
        if _pending_passes(ws, r, amounts, changed, cfg):
            to_process_rows.append(r)
    total_to_process = len(to_process_rows)
    if cfg.limit is not None:
        total_to_process = min(total_to_process, cfg.limit)

    hooks.event("start", total_rows=len(rows), to_process=total_to_process)

    use_checkpoint = cfg.checkpoint and total_to_process >= CHECKPOINT_MIN_ROWS
    checkpoint_after = max(1, math.ceil(CHECKPOINT_FRACTION * total_to_process)) if use_checkpoint else None

    # ---- browser --------------------------------------------------------
    assert_not_locked(xlsx)
    pw, browser, page = connect_page(cfg.cdp_url)
    locked_url = page.url
    hooks.log(f"Attached to tab: {locked_url}\n")

    backup_path = None
    if not cfg.no_backup:
        backup_path, fresh = make_backup(xlsx)
        hooks.log(f"Backup: {backup_path}"
                  + ("" if fresh else "  (from a previous unfinished run)") + "\n")

    summary = dict(processed=0, already_done=0, skipped=0, messages=0,
                   failed_rows=[], aborted=None, stopped=False, run_ok=False)
    dirty = False
    saved_ok = False
    checkpoint_passed = not use_checkpoint

    def save_all():
        nonlocal saved_ok
        wb.save(xlsx)
        save_state(xlsx, state)
        saved_ok = True
        hooks.event("saved", path=os.path.abspath(xlsx))

    def set_globals():
        for key, val in (("closing_date", closing_date),
                         ("underwriter", underwriter), ("state", state_val)):
            if val == "":
                continue
            try:
                set_text(page, key, val)
            except LookupError as e:
                hooks.log(f"WARNING: {e}")

    try:
        set_globals()

        for rownum in rows:
            if hooks.should_stop():
                summary["stopped"] = True
                break

            stage = "starting row"
            try:
                if page.url != locked_url:
                    raise RuntimeError(f"the open tab is no longer the Testing "
                                       f"Calculator (it is now {page.url})")

                inp = read_row_inputs(ws, rownum)
                amounts = [(c, amount_for[c], PASS_COLUMNS[c][1])
                           for c in PASS_COLUMNS if not is_blank(inp["expected"][c])]
                reason = skip_reason(inp, amounts)
                sig = row_signature(inp, gsig)
                stored = state["rows"].get(str(rownum))
                inputs_changed = stored is not None and stored.get("sig") != sig

                if reason:
                    ws[f"K{rownum}"].value = reason
                    state["rows"][str(rownum)] = {"sig": sig, "cols": ["K"]}
                    dirty = True
                    summary["skipped"] += 1
                    hooks.event("row_skip", row=rownum, reason=reason)
                    continue

                pending = _pending_passes(ws, rownum, amounts, inputs_changed, cfg)
                if not pending:
                    state["rows"][str(rownum)] = {
                        "sig": sig, "cols": _filled_cols(ws, rownum, amounts)}
                    summary["already_done"] += 1
                    hooks.event("row_done_already", row=rownum, quick_ref=inp["quick_ref"])
                    continue

                if inputs_changed:
                    hooks.event("row_changed", row=rownum)
                    for (_c, _a, kcol) in pending:
                        ws[f"{kcol}{rownum}"].value = None

                hooks.event("row_begin", row=rownum, quick_ref=inp["quick_ref"],
                            liability=inp["liability"], coverage=inp["coverage"])

                stage = "entering Fee Item Code"
                set_text(page, "fee_item_code", inp["quick_ref"])
                page.wait_for_timeout(500)
                stage = "selecting Liability Type"
                set_select(page, "liability_type", inp["liability"])
                stage = "entering Coverage Type Code"
                set_text(page, "coverage_type_code", inp["coverage"])

                if inp["question"]:
                    stage = "entering Boolean Question / Answer"
                    set_text(page, "boolean_question", inp["question"])
                    if inp["answer"]:
                        set_select(page, "boolean_answer", inp["answer"])
                else:
                    try:
                        set_text(page, "boolean_question", "")
                    except LookupError:
                        pass

                fee_loc = field(page, "fee_amount")
                err_loc = field(page, "error_msg")
                done_cols = []

                for col, amt, kcol in pending:
                    stage = f"amount pass {col} -> {kcol}{rownum} (amount {amt})"
                    before_fee = read_value(fee_loc)
                    before_err = read_value(err_loc)
                    set_text(page, "amount", amt)

                    wait_not_loading(page)
                    field(page, "calculate_button").click()
                    page.wait_for_timeout(200)
                    wait_not_loading(page)

                    deadline = time.time() + cfg.settle_ms / 1000.0
                    while time.time() < deadline:
                        page.wait_for_timeout(200)
                        if (read_value(fee_loc) != before_fee
                                or read_value(err_loc) != before_err):
                            break
                    page.wait_for_timeout(300)

                    fee = read_value(fee_loc)
                    err = read_value(err_loc)
                    fee_num = parse_money(fee)

                    cell = ws[f"{kcol}{rownum}"]
                    if fee_num is not None and fee_num != 0:
                        cell.value = fee_num
                        cell.number_format = '"$"#,##0.00'
                        value, is_msg = fee_num, False
                    else:
                        cell.value = err or fee
                        value, is_msg = (err or fee), bool(err)

                    dirty = True
                    done_cols.append(kcol)
                    summary["messages"] += is_msg
                    hooks.event("pass_done", row=rownum, col=col, kcol=kcol,
                                amount=amt, value=value, is_message=is_msg)

                state["rows"][str(rownum)] = {"sig": sig, "cols": done_cols}
                summary["processed"] += 1
                hooks.event("progress", done=summary["processed"],
                            total=total_to_process)

                if summary["processed"] % SAVE_EVERY_ROWS == 0:
                    save_all()

                if cfg.limit is not None and summary["processed"] >= cfg.limit:
                    hooks.log(f"\nReached limit {cfg.limit}; stopping.")
                    break

                if not checkpoint_passed and summary["processed"] >= checkpoint_after:
                    checkpoint_passed = True
                    save_all()
                    hooks.event("checkpoint", done=summary["processed"],
                                total=total_to_process, path=os.path.abspath(xlsx))
                    if hooks.wait_at_checkpoint() == "stop":
                        summary["stopped"] = True
                        break
                    # reload - the user may have edited the sheet during the pause
                    wb = load_workbook(xlsx)
                    ws = wb[cfg.sheet]
                    state = load_state(xlsx)
                    closing_date, underwriter, state_val, amount_for = read_globals(ws)
                    new_gsig = globals_signature(closing_date, underwriter, state_val)
                    if new_gsig != gsig:
                        gsig = new_gsig
                        set_globals()
                    hooks.event("resumed")

            except KeyboardInterrupt:
                raise
            except Exception as e:  # noqa: BLE001
                friendly = explain_error(e)
                if cfg.keep_going:
                    summary["failed_rows"].append(rownum)
                    hooks.event("row_failed", row=rownum, stage=stage, friendly=friendly)
                    continue
                summary["aborted"] = (rownum, stage, friendly)
                break

        if summary["aborted"] is None and not summary["stopped"]:
            summary["run_ok"] = True
    finally:
        try:
            if dirty:
                save_all()
            else:
                hooks.log("\nNothing to write.")
        except Exception as e:  # noqa: BLE001
            hooks.log(f"\nERROR saving workbook: {explain_error(e)}")

        if backup_path and os.path.exists(backup_path):
            if summary["run_ok"] and (saved_ok or not dirty):
                try:
                    os.remove(backup_path)
                    hooks.log("Run finished OK - backup deleted.")
                except OSError:
                    pass
            else:
                hooks.log(f"Backup kept (run did not finish clean): {backup_path}")

        try:
            browser.close()   # CDP: detaches only, your Edge stays open
        except Exception:
            pass
        pw.stop()

    if summary["aborted"]:
        r, stg, fr = summary["aborted"]
        hooks.event("aborted", row=r, stage=stg, friendly=fr)
    elif summary["stopped"]:
        hooks.event("stopped")
    hooks.event("finished", summary=summary)

    # Only after a run that actually reached the end: ask whether to clean up the
    # resume/state file.  Keep it if the user isn't satisfied so they can edit
    # rows and rerun with change-detection intact.
    if summary["run_ok"]:
        sp = state_path_for(xlsx)
        if hooks.ask_satisfied(summary):
            try:
                os.remove(sp)
            except OSError:
                pass
            hooks.event("state_deleted")
        elif os.path.exists(sp):
            hooks.event("state_kept")

    return summary


# --------------------------------------------------------------------- CLI -----

def run_dry(cfg: EngineConfig, hooks: Hooks) -> None:
    from openpyxl import load_workbook

    if not os.path.isfile(cfg.xlsx):
        raise SetupError(f"file not found: {cfg.xlsx}")
    wb = load_workbook(cfg.xlsx)
    if cfg.sheet not in wb.sheetnames:
        raise SetupError(f"no sheet '{cfg.sheet}'. Sheets: {wb.sheetnames}")
    ws = wb[cfg.sheet]
    state = load_state(cfg.xlsx)
    closing_date, underwriter, state_val, amount_for = read_globals(ws)
    gsig = globals_signature(closing_date, underwriter, state_val)

    hooks.log(f"Globals: State={state_val!r} Underwriter={underwriter!r} "
              f"ClosingDate={closing_date!r}\n")
    for r in data_row_numbers(ws, cfg.start_row):
        inp = read_row_inputs(ws, r)
        amounts = [(c, amount_for[c], PASS_COLUMNS[c][1])
                   for c in PASS_COLUMNS if not is_blank(inp["expected"][c])]
        reason = skip_reason(inp, amounts)
        if reason:
            hooks.log(f"row {r}: {reason}  -> K{r}")
            continue
        stored = state["rows"].get(str(r))
        changed = stored is not None and stored.get("sig") != row_signature(inp, gsig)
        pending = _pending_passes(ws, r, amounts, changed, cfg)
        if not pending:
            hooks.log(f"row {r}: {inp['quick_ref']}  already filled - would skip")
            continue
        tag = "  [inputs changed]" if changed else ""
        passes = ", ".join(f"{c}:{amt}->{kc}{r}" for c, amt, kc in pending)
        hooks.log(f"row {r}: {inp['quick_ref']}  {inp['liability']} | "
                  f"{inp['coverage']}  passes[{passes}]{tag}")
    hooks.log("\nDry run only - nothing written, browser not touched.")


def parse_args(argv):
    p = argparse.ArgumentParser(description="Fill the STRAFE Testing Calculator from Excel.")
    p.add_argument("xlsx", help="path to the test workbook (.xlsx)")
    p.add_argument("--sheet", default=SHEET_NAME_DEFAULT)
    p.add_argument("--limit", type=int, default=None)
    p.add_argument("--start-row", type=int, default=FIRST_DATA_ROW, dest="start_row")
    p.add_argument("--cdp-url", default=CDP_URL_DEFAULT, dest="cdp_url")
    p.add_argument("--settle-ms", type=int, default=4000, dest="settle_ms")
    p.add_argument("--dry-run", action="store_true")
    p.add_argument("--overwrite", action="store_true",
                   help="redo every row even if its Actual Value cells are filled")
    p.add_argument("--retry-messages", action="store_true", dest="retry_messages",
                   help="also redo cells that hold a text message instead of a $ amount")
    p.add_argument("--keep-going", action="store_true", dest="keep_going",
                   help="on a row error, log it and continue instead of stopping")
    p.add_argument("--no-checkpoint", action="store_true", dest="no_checkpoint",
                   help="do not pause at ~10%%")
    p.add_argument("--no-backup", action="store_true", dest="no_backup",
                   help="do not make the '(backup)' safety copy")
    p.add_argument("--list-fields", action="store_true", dest="list_fields")
    return p.parse_args(argv)


def main(argv=None):
    args = parse_args(argv)
    hooks = ConsoleHooks()

    if args.list_fields:
        pw, browser, page = connect_page(args.cdp_url)
        try:
            list_fields(page)
        finally:
            try:
                browser.close()
            except Exception:
                pass
            pw.stop()
        return

    cfg = EngineConfig(
        xlsx=args.xlsx, sheet=args.sheet, start_row=args.start_row, limit=args.limit,
        cdp_url=args.cdp_url, settle_ms=args.settle_ms, overwrite=args.overwrite,
        retry_messages=args.retry_messages, keep_going=args.keep_going,
        checkpoint=not args.no_checkpoint, no_backup=args.no_backup,
    )

    if args.dry_run:
        run_dry(cfg, hooks)
        return

    summary = run_engine(cfg, hooks)
    if summary["aborted"]:
        sys.exit(1)


if __name__ == "__main__":
    try:
        main()
    except KeyboardInterrupt:
        print("\nInterrupted (Ctrl+C). Progress up to the last saved point is kept; "
              "run the same command to continue.", flush=True)
        sys.exit(130)
    except SetupError as e:
        print(f"\nERROR: {e}", flush=True)
        sys.exit(2)
    except SystemExit:
        raise
    except Exception as e:  # noqa: BLE001
        print(f"\nERROR: {explain_error(e)}", flush=True)
        sys.exit(1)